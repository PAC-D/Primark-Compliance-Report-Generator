from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def _make_thin_border():
    thin = Side(style='thin')
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _month_abbr(month: int) -> str:
    import calendar
    return calendar.month_abbr[month]


def _build_yearly_performance_sheet(wb, supplier_trend, month_labels, year):
    ws = wb.create_sheet("Yearly Performance")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    ws["A1"] = f"Suppliers Yearly Performance – January {year} to {chart_end_caption(month_labels, year)}"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    num_months = len(month_labels)

    ws.merge_cells(f"C5:{get_column_letter(2 + num_months)}5")
    ws["C5"] = str(year)
    ws["C5"].font = Font(bold=True, color="FFFFFF", size=11)
    ws["C5"].fill = PatternFill(start_color="4A5568", end_color="4A5568", fill_type="solid")
    ws["C5"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 22

    headers = ["#", "Supplier Name"] + month_labels + ["Total"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill(start_color="1B2838", end_color="1B2838", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _make_thin_border()

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 45
    for i in range(num_months):
        ws.column_dimensions[get_column_letter(3 + i)].width = 10
    total_col = 3 + num_months
    ws.column_dimensions[get_column_letter(total_col)].width = 12

    for idx, supplier in enumerate(supplier_trend, start=1):
        row = 6 + idx
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=supplier["supplier"])

        for i, rate in enumerate(supplier["rates"]):
            cell = ws.cell(row=row, column=3 + i)
            if rate is not None:
                cell.value = rate / 100.0
                cell.number_format = "0.00%"
                if rate < 40:
                    cell.font = Font(color="DC2626", bold=True)
                elif rate < 60:
                    cell.font = Font(color="EA580C")
            else:
                cell.value = ""
            cell.alignment = Alignment(horizontal="center")
            cell.border = _make_thin_border()
            if idx % 2 == 0:
                cell.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

        last_col_letter = get_column_letter(2 + num_months)
        total_cell = ws.cell(row=row, column=total_col)
        total_cell.value = f"=IFERROR(AVERAGE(C{row}:{last_col_letter}{row}),\"\")"
        total_cell.number_format = "0.00%"
        total_cell.alignment = Alignment(horizontal="center")
        total_cell.border = _make_thin_border()
        total_cell.font = Font(bold=True)
        if idx % 2 == 0:
            total_cell.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

        for col in (1, 2):
            cell = ws.cell(row=row, column=col)
            cell.border = _make_thin_border()
            cell.alignment = Alignment(horizontal="center" if col == 1 else "left")
            if idx % 2 == 0:
                cell.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    ws.freeze_panes = "C7"


def chart_end_caption(month_labels, year):
    if not month_labels:
        return f"{year}"
    last_label = month_labels[-1]
    try:
        m_abbr, yy = last_label.split(" '")
        mm = next((i for i, name in enumerate(
            ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
        ) if name == m_abbr), None)
        import calendar
        return f"{calendar.month_name[mm]} 20{int(yy)}"
    except Exception:
        return f"{year}"
