import tkinter as tk
from tkinter import filedialog
import os
from pathlib import Path
from datetime import date

from processor import is_file_processed, process_file, copy_to_data
from database import get_session, text, clear_supplier_emails, clear_raw_data, clear_all_tables, insert_supplier_email, init_supplier_emails_table
from renderer import render_report
from summary_renderer import render_summary_report
from summary_excel_renderer import render_summary_excel
from kpi_renderer import render_kpi_report
from metrics import get_supplier_kpis
from pdfgen import html_to_pdf, html_to_pdf_mixed, _merge_pdfs, _add_page_numbers
import sys
sys.path.insert(0, str(Path(__file__).parent.parent / "supplier-data-extractor"))
from extract_supplier_data import get_month_year_input, extract_supplier_data
from safe_filename import safe_filename
from logging_config import setup_logging, get_logger
import yaml
import openpyxl

setup_logging()
logger = get_logger(__name__)

CN_REGION_COUNTRIES = {"China", "Myanmar", "Cambodia", "Vietnam", "VietNam", "Indonesia"}

BD_REGION_COUNTRIES = {"Bangladesh", "India", "Pakistan", "Sri Lanka", "Morocco", "Turkey"}


def _get_report_region(origin_country: str) -> str:
    """Returns 'CN Region' if origin_country contains any CN region country, else 'BD Region'."""
    countries = {c.strip() for c in origin_country.split(",")}
    return "CN Region" if countries & CN_REGION_COUNTRIES else "BD Region"

MAX_LOG_LINES = 50


def print_banner():
    banner = """
============================================
   TPS Compliance Report Generator
============================================
   [1] Process New File
   [2] Generate Reports
   [3] Generate Summary Report
   [4] View Logs
   [5] Extract Supplier Data
   [6] Import Supplier Emails
   [7] Generate KPI Report
   [8] Clear Database
   [0] Exit
============================================
    """
    print(banner)


def handle_process_file():
    root = tk.Tk()
    root.withdraw()
    try:
        filepaths = filedialog.askopenfilenames(
            title="Select Transaction Data Excel File(s)",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
        )
    finally:
        root.destroy()

    if not filepaths:
        return

    for filepath in filepaths:
        filename = os.path.basename(filepath)
        logger.info(f"Processing file: {filepath}")

        if is_file_processed(filepath):
            logger.info("File already processed, skipping")
            print(f"File already processed: {filename}")
            response = input(f"Re-process {filename}? (y/n): ").strip().lower()
            if response != "y":
                continue

        try:
            row_count = process_file(filepath)
            copy_to_data(filepath)
            logger.info(f"Inserted {row_count} rows from {filename}")
            logger.info("File copied to data/")
            print(f"Successfully processed: {filename} ({row_count} rows)")
        except Exception as e:
            logger.warning(f"Error processing file {filename}: {e}")
            print(f"Warning: {filename} failed - {e}. Skipping.")
            continue


def handle_generate_reports():
    config_path = Path(__file__).parent / "config.yaml"
    try:
        with open(config_path) as f:
            config = yaml.safe_load(f)
    except Exception as e:
        print(f"Error loading config: {e}")
        return
    output_dir = Path(__file__).parent / config["paths"]["output_dir"]
    output_dir.mkdir(exist_ok=True)

    print("\n=== Report Month Selection ===")
    print("Enter year and month for the report (or press Enter for default: previous month)")
    year_input = input("Year (e.g., 2026) [default: previous month]: ").strip()
    month_input = input("Month (1-12) [default: previous month]: ").strip()

    report_month = None
    if month_input:
        try:
            year = int(year_input) if year_input else (date.today().year)
            month = int(month_input)
            report_month = date(year, month, 1)
            print(f"Using specified month: {report_month.strftime('%B %Y')}")
        except ValueError:
            print("Invalid date. Using default.")
            report_month = None

    if report_month is None:
        print("Using default report month.")
        today = date.today()
        year = today.year if today.month > 1 else today.year - 1
        month = today.month - 1 if today.month > 1 else 12
        report_month = date(year, month, 1)

    session = get_session()
    try:
        result = session.execute(
            text("SELECT DISTINCT supplier FROM raw_data WHERE EXTRACT(YEAR FROM handover_date) = :year AND EXTRACT(MONTH FROM handover_date) = :month"),
            {"year": report_month.year, "month": report_month.month}
        )
        suppliers = [row[0] for row in result]
    finally:
        session.close()

    if not suppliers:
        print("No suppliers found in database.")
        return

    logger.info(f"Starting report generation for {len(suppliers)} suppliers")
    success_count = 0
    skip_count = 0

    for i, supplier in enumerate(suppliers, 1):
        logger.info(f"Generating report {i}/{len(suppliers)} for: {supplier}")
        print(f"Generating report {i}/{len(suppliers)} for: {supplier}...")
        try:
            kpis = get_supplier_kpis(supplier, report_month)
            origin_country = kpis.get("origin_country", "")
            region = _get_report_region(origin_country)

            portrait_html, landscape_html = render_report(supplier, report_month)

            safe_name = safe_filename(supplier)
            file_date = report_month.strftime("%B_%Y") if report_month else date.today().strftime("%B_%Y")
            month_folder = output_dir / f"Report_{file_date}"
            region_folder = month_folder / region
            region_folder.mkdir(parents=True, exist_ok=True)
            output_path = region_folder / f"{safe_name}_{file_date}.pdf"

            # Remove existing file to avoid conflicts
            if output_path.exists():
                output_path.unlink()

            html_to_pdf_mixed(portrait_html, landscape_html, str(output_path))
            logger.info(f"Report saved: {output_path}")
            success_count += 1
        except Exception as e:
            logger.error(f"Error generating report for {supplier}: {e}")
            logger.info(f"Skipped: {supplier} - {e}")
            print(f"  Skipped: {supplier} - {e}")
            skip_count += 1

    logger.info(f"Generated {success_count} reports, {skip_count} skipped")
    print(f"Generated {success_count} reports, {skip_count} skipped.")


def handle_view_logs():
    logs_dir = Path(__file__).parent / "logs"
    log_file = logs_dir / "run.log"

    try:
        with open(log_file) as f:
            lines = f.readlines()
        for line in lines[-MAX_LOG_LINES:]:
            print(line.rstrip())
    except FileNotFoundError:
        print("No logs found yet.")
    except Exception as e:
        print(f"Error reading logs: {e}")


def handle_generate_summary_report():
    config_path = Path(__file__).parent / "config.yaml"
    try:
        with open(config_path) as f:
            config = yaml.safe_load(f)
    except Exception as e:
        print(f"Error loading config: {e}")
        return
    output_dir = Path(__file__).parent / config["paths"]["output_dir"]
    output_dir.mkdir(exist_ok=True)

    print("\n=== Summary Report Month Selection ===")
    print("Enter year and month for the report (or press Enter for default: previous month)")
    year_input = input("Year (e.g., 2026) [default: previous month]: ").strip()
    month_input = input("Month (1-12) [default: previous month]: ").strip()

    report_month = None
    if month_input:
        try:
            year = int(year_input) if year_input else (date.today().year)
            month = int(month_input)
            report_month = date(year, month, 1)
            print(f"Using specified month: {report_month.strftime('%B %Y')}")
        except ValueError:
            print("Invalid date. Using default.")
            report_month = None

    if report_month is None:
        print("Using default report month.")
        today = date.today()
        year = today.year if today.month > 1 else today.year - 1
        month = today.month - 1 if today.month > 1 else 12
        report_month = date(year, month, 1)

    print("Rendering summary report...")
    try:
        wb = render_summary_excel(report_month)

        file_date = report_month.strftime("%B_%Y") if report_month else date.today().strftime("%B_%Y")
        output_path = output_dir / f"Summary_Report_{file_date}.xlsx"

        wb.save(str(output_path))

        logger.info(f"Summary report saved: {output_path}")
        print(f"Summary report generated: {output_path}")
    except Exception as e:
        logger.error(f"Error generating summary report: {e}")
        print(f"Error generating summary report: {e}")


def handle_extract_supplier_data():
    try:
        year, month = get_month_year_input()
        print(f"\nExtracting data for {month}/{year}...")
        extract_supplier_data(year, month)
    except Exception as e:
        logger.error(f"Error extracting supplier data: {e}")
        print(f"Error: {e}")


def handle_import_emails():
    root = tk.Tk()
    root.withdraw()
    try:
        filepath = filedialog.askopenfilename(
            title="Select Supplier Emails Excel File",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
        )
    finally:
        root.destroy()

    if not filepath:
        return

    logger.info(f"Importing supplier emails from: {filepath}")

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active

        header_row = [cell.value for cell in ws[1]]
        if len(header_row) < 2:
            print("Error: Excel must have at least 2 columns (Supplier Name, Email)")
            return

        supplier_col = None
        email_col = None
        for i, h in enumerate(header_row):
            h_lower = (str(h) if h is not None else "").strip().lower()
            if "supplier" in h_lower and "name" in h_lower:
                supplier_col = i + 1
            elif "email" in h_lower:
                email_col = i + 1

        if not supplier_col or not email_col:
            print("Error: Could not find 'Supplier Name' and 'Email' columns")
            return

        init_supplier_emails_table()
        clear_supplier_emails()

        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            supplier_name = row[supplier_col - 1]
            email = row[email_col - 1]
            if supplier_name and email:
                for single_email in str(email).replace(";", ",").split(","):
                    single_email = single_email.strip()
                    if single_email:
                        insert_supplier_email(str(supplier_name), single_email)
                        count += 1

        logger.info(f"Imported {count} supplier emails")
        print(f"Successfully imported {count} supplier emails")

    except Exception as e:
        logger.error(f"Error importing supplier emails: {e}")
        print(f"Error importing supplier emails: {e}")


def handle_generate_kpi_report():
    config_path = Path(__file__).parent / "config.yaml"
    try:
        with open(config_path) as f:
            config = yaml.safe_load(f)
    except Exception as e:
        print(f"Error loading config: {e}")
        return
    output_dir = Path(__file__).parent / config["paths"]["output_dir"]
    output_dir.mkdir(exist_ok=True)

    print("\n=== KPI Report Month Selection ===")
    print("Enter year and month for the report (or press Enter for default: previous month)")
    year_input = input("Year (e.g., 2026) [default: previous month]: ").strip()
    month_input = input("Month (1-12) [default: previous month]: ").strip()

    report_month = None
    if month_input:
        try:
            year = int(year_input) if year_input else date.today().year
            month = int(month_input)
            report_month = date(year, month, 1)
            print(f"Using specified month: {report_month.strftime('%B %Y')}")
        except ValueError:
            print("Invalid date. Using default.")
            report_month = None

    if report_month is None:
        print("Using default report month.")
        today = date.today()
        year = today.year if today.month > 1 else today.year - 1
        month = today.month - 1 if today.month > 1 else 12
        report_month = date(year, month, 1)

    print("Rendering KPI report...")
    try:
        html = render_kpi_report(report_month)

        file_date = report_month.strftime("%B_%Y")
        output_path = output_dir / f"KPI_Report_{file_date}.pdf"

        if output_path.exists():
            output_path.unlink()

        html_to_pdf(html, str(output_path))
        logger.info(f"KPI report saved: {output_path}")
        print(f"KPI report generated: {output_path}")
    except ValueError as e:
        logger.warning(f"KPI report skipped: {e}")
        print(f"KPI report skipped: {e}")
    except Exception as e:
        logger.error(f"Error generating KPI report: {e}")
        print(f"Error generating KPI report: {e}")


def handle_reload_data():
    data_dir = Path(__file__).parent / "data"
    filepaths = sorted(data_dir.glob("*.xlsx")) + sorted(data_dir.glob("*.xls"))

    if not filepaths:
        print("No files found in data/ folder.")
        return

    for filepath in filepaths:
        filename = filepath.name
        logger.info(f"Reloading file: {filepath}")
        try:
            row_count = process_file(str(filepath))
            logger.info(f"Inserted {row_count} rows from {filename}")
            print(f"Successfully processed: {filename} ({row_count} rows)")
        except Exception as e:
            logger.warning(f"Error processing file {filename}: {e}")
            print(f"Warning: {filename} failed - {e}. Skipping.")
            continue


def handle_clear_database():
    print("\n=== Clear Database ===")
    response = input("This will DELETE ALL rows from raw_data and supplier_emails (tables are kept). Continue? (y/n): ").strip().lower()
    if response != "y":
        print("Canceled.")
        return
    try:
        clear_all_tables()
        logger.info("Database cleared: all rows removed from raw_data and supplier_emails")
        print("Database cleared: all rows removed. Tables preserved.")
    except Exception as e:
        logger.error(f"Error clearing database: {e}")
        print(f"Error clearing database: {e}")
        return

    reload_response = input("Reload data from the data/ folder? (y/n): ").strip().lower()
    if reload_response == "y":
        handle_reload_data()


def main():
    print_banner()
    while True:
        choice = input("\nSelect option: ").strip()
        if choice == "1":
            handle_process_file()
        elif choice == "2":
            handle_generate_reports()
        elif choice == "3":
            handle_generate_summary_report()
        elif choice == "4":
            handle_view_logs()
        elif choice == "5":
            handle_extract_supplier_data()
        elif choice == "6":
            handle_import_emails()
        elif choice == "7":
            handle_generate_kpi_report()
        elif choice == "8":
            handle_clear_database()
        elif choice == "0":
            print("Goodbye!")
            break
        else:
            print("Invalid option. Try again.")


if __name__ == "__main__":
    main()