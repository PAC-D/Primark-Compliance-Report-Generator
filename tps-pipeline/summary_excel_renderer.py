from datetime import date
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from _suppliers_below import build_follow_up_required_sheet as _build_suppliers_below_threshold_sheet
from _supplier_trend import _build_yearly_performance_sheet
from metrics import (
    get_report_month,
    get_all_suppliers_24month_trend,
    get_summary_sources,
    get_summary_country_compliance,
    get_summary_country_monthly_trend,
    get_summary_line_details,
    get_top_suppliers,
    get_suppliers_below_threshold,
    get_supplier_emails_map,
)


def _make_thin_border():
    thin = Side(style='thin')
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def render_summary_excel(report_month: date | None = None) -> Workbook:
    if report_month is None:
        report_month = get_report_month()

    wb = Workbook()
    wb.remove(wb.active)

    sources = get_summary_sources(report_month)
    compliance_data = get_summary_country_compliance(report_month)
    country_data, chart_labels, country_order = get_summary_country_monthly_trend(report_month)
    line_details, line_totals = get_summary_line_details(report_month)
    top_suppliers = get_top_suppliers(report_month)
    suppliers_below = get_suppliers_below_threshold(report_month)
    all_supplier_trend, month_labels = get_all_suppliers_24month_trend(report_month)
    email_map = get_supplier_emails_map()

    month_name = report_month.strftime("%B")
    year = report_month.year

    _build_summary_sheet(wb, sources, compliance_data, country_data,
                         chart_labels, country_order, month_name, year)
    _build_visibility_sheet(wb, line_details, line_totals, month_name, year)
    _build_top_performers_sheet(wb, top_suppliers, month_name, year)
    _build_suppliers_below_threshold_sheet(wb, suppliers_below, email_map)
    _build_yearly_performance_sheet(wb, all_supplier_trend, month_labels, year)

    return wb


def _add_trend_chart(ws, country_data, chart_labels, country_order, start_row):
    num_countries = len(country_order)
    num_labels = len(chart_labels)

    label_col = 1
    val_start_col = 2
    val_end_col = val_start_col + num_countries - 1

    for j, country in enumerate(country_order, val_start_col):
        cell = ws.cell(row=start_row, column=j, value=country)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill(start_color="1B2838", end_color="1B2838", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
        cell.border = _make_thin_border()

    for i, label in enumerate(chart_labels):
        row = start_row + 1 + i
        ws.cell(row=row, column=label_col, value=label)
        ws.cell(row=row, column=label_col).font = Font(bold=True, size=9)
        ws.cell(row=row, column=label_col).alignment = Alignment(horizontal="center")
        for j, country in enumerate(country_order, val_start_col):
            vals = country_data.get(country, [])
            v = vals[i] if i < len(vals) else 0
            ws.cell(row=row, column=j, value=v)

    data = Reference(ws, min_col=val_start_col, min_row=start_row,
                     max_col=val_end_col, max_row=start_row + num_labels)
    cats = Reference(ws, min_col=label_col, min_row=start_row + 1, max_row=start_row + num_labels)

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "Compliance Rate by Country (%)"
    chart.y_axis.title = "Compliance Rate (%)"
    chart.x_axis.title = "Month"
    chart.style = 10
    chart.width = 28
    chart.height = 14

    chart.add_data(data, from_rows=False, titles_from_data=True)
    chart.set_categories(cats)

    ws.add_chart(chart, f"A{start_row + num_labels + 2}")


def _build_summary_sheet(wb, sources, compliance_data, country_data,
                         chart_labels, country_order, month_name, year):
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:F1")
    ws["A1"] = f"TPS COMPLIANCE REPORT – Summary for {month_name} {year}"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 25

    ws["A3"] = "Sources"
    ws["A3"].font = Font(bold=True, size=11)

    ws["A4"] = "Supplier"
    ws["B4"] = "Factory"
    for cell in ws["A4:B4"][0]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        cell.alignment = Alignment(horizontal="right")
        cell.border = _make_thin_border()

    ws["A5"] = sources["total_suppliers"]
    ws["B5"] = sources["total_factories"]
    ws["A5"].alignment = Alignment(horizontal="right")
    ws["B5"].alignment = Alignment(horizontal="right")
    ws["A5"].number_format = "#,##0"
    ws["B5"].number_format = "#,##0"

    ws["A7"] = "Volume (Compliant Only)"
    ws["A7"].font = Font(bold=True, size=11)

    ws["A8"] = "Country"
    ws["B8"] = "Compliance Score"
    for cell in ws["A8:B8"][0]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        cell.alignment = Alignment(horizontal="right" if cell.column == 2 else "left")
        cell.border = _make_thin_border()

    for i, item in enumerate(compliance_data):
        row = 9 + i
        ws.cell(row=row, column=1, value=item["country"])
        rate = item["compliance_rate"] / 100.0
        ws.cell(row=row, column=2, value=rate)
        ws.cell(row=row, column=2).number_format = "0.00%"
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="right")

    avg_row = 9 + len(compliance_data)
    ws.cell(row=avg_row, column=1, value="Average")
    preset_rates = "B9:B" + str(8 + len(compliance_data))
    ws.cell(row=avg_row, column=2, value=f"=AVERAGE({preset_rates})")
    ws.cell(row=avg_row, column=2).number_format = "0.00%"
    ws.cell(row=avg_row, column=1).font = Font(bold=True)
    ws.cell(row=avg_row, column=2).font = Font(bold=True)
    ws.cell(row=avg_row, column=2).alignment = Alignment(horizontal="right")

    chart_start_row = avg_row + 3
    _add_trend_chart(ws, country_data, chart_labels, country_order, chart_start_row)

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18


def _build_visibility_sheet(wb, line_details, line_totals, month_name, year):
    ws = wb.create_sheet("Visibility")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:E1")
    ws["A1"] = f"Visibility – {month_name} {year}"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    headers = ["Origin Country", "Destination PO", "Reflected Lines", "Approved", "Rejected"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        cell.alignment = Alignment(horizontal="right" if col > 1 else "left")
        cell.border = _make_thin_border()

    for i, item in enumerate(line_details):
        row = 3 + i
        ws.cell(row=row, column=1, value=item["country"])
        ws.cell(row=row, column=2, value=item["dest_po_count"])
        ws.cell(row=row, column=3, value=item["reflected_lines"])
        ws.cell(row=row, column=4, value=item["approved"])
        ws.cell(row=row, column=5, value=item["rejected"])
        for col in range(2, 6):
            ws.cell(row=row, column=col).alignment = Alignment(horizontal="right")

    total_row = 3 + len(line_details)
    ws.cell(row=total_row, column=1, value="Total")
    ws.cell(row=total_row, column=1).font = Font(bold=True)
    ws.cell(row=total_row, column=2, value=line_totals["dest_po_count"])
    ws.cell(row=total_row, column=3, value=line_totals["reflected_lines"])
    ws.cell(row=total_row, column=4, value=line_totals["approved"])
    ws.cell(row=total_row, column=5, value=line_totals["rejected"])
    for col in range(1, 6):
        cell = ws.cell(row=total_row, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="E8F4FD", end_color="E8F4FD", fill_type="solid")
        cell.border = _make_thin_border()
        if col > 1:
            cell.alignment = Alignment(horizontal="right")

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.freeze_panes = "A3"


def _build_top_performers_sheet(wb, top_suppliers, month_name, year):
    ws = wb.create_sheet("Top Performers")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:G1")
    ws["A1"] = f"Top Suppliers  – {month_name} {year}"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    headers = ["#", "Origin", "Supplier", "No. of Factories", "Destination PO", "Reflected Lines", "Compliance Score"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        cell.alignment = Alignment(horizontal="right" if col in (1, 4, 5, 6, 7) else "left")
        cell.border = _make_thin_border()

    for i, supplier in enumerate(top_suppliers):
        row = 3 + i
        ws.cell(row=row, column=1, value=supplier["sl"])
        ws.cell(row=row, column=2, value=supplier["origin_countries"])
        ws.cell(row=row, column=3, value=supplier["supplier"])
        ws.cell(row=row, column=4, value=supplier["factory_count"])
        ws.cell(row=row, column=5, value=supplier["dest_po_count"])
        ws.cell(row=row, column=6, value=supplier["reflected_lines"])
        rate = supplier["compliance_rate"] / 100.0
        ws.cell(row=row, column=7, value=rate)
        ws.cell(row=row, column=7).number_format = "0.00%"
        ws.cell(row=row, column=7).alignment = Alignment(horizontal="right")

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 18
    ws.freeze_panes = "A3"
