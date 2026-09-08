from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def _make_thin_border():
    thin = Side(style='thin')
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def build_follow_up_required_sheet(wb, suppliers_below, email_map=None):
    if email_map is None:
        email_map = {}
    ws = wb.create_sheet("Follow-Up Required")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    title = "Follow-Up Required (Avg. Compliance <40% – Last 3 Months)"
    if suppliers_below:
        first = suppliers_below[0]
        if "month_label" in first:
            title = f"{title} – {first['month_label']}"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    headers = ["Sl ", "Origin Country ", "Supplier Name ", "No of Factories ",
               "Destination PO Qty ", "Reflected Lines ", "Compliance Score ", "Email "]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill(start_color="156082", end_color="156082", fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if col == 1 else "left")
        cell.border = _make_thin_border()

    for i, supplier in enumerate(suppliers_below):
        row = 3 + i
        ws.cell(row=row, column=1, value=supplier["sl"])
        ws.cell(row=row, column=2, value=supplier["origin_country"])
        ws.cell(row=row, column=3, value=supplier["supplier"])
        ws.cell(row=row, column=4, value=supplier["factory_count"])
        ws.cell(row=row, column=5, value=supplier["dest_po_count"])
        ws.cell(row=row, column=6, value=supplier["reflected_lines"])

        rate = supplier.get("avg_3month")
        if rate is None:
            rate = supplier.get("compliance_rate", 0)
        ws.cell(row=row, column=7, value=rate / 100.0)
        ws.cell(row=row, column=7).number_format = "0.0%"
        if rate < 30:
            ws.cell(row=row, column=7).font = Font(color="DC2626", bold=True)
        elif rate < 40:
            ws.cell(row=row, column=7).font = Font(color="EA580C", bold=True)

        email = email_map.get(supplier["supplier"].strip().lower(), "")
        ws.cell(row=row, column=8, value=email)

        ws.row_dimensions[row].height = 15

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 40
    ws.freeze_panes = "A3"


# Kept for backward-compat with any external callers; routes to new name+layout.
def _build_suppliers_below_threshold_sheet(*args, **kwargs):
    return build_follow_up_required_sheet(*args, **kwargs)
